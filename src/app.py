from datetime import datetime

from flask import Flask, render_template, request, redirect, url_for, flash, jsonify, make_response, send_file
from flask_mysqldb import MySQL
from flask_wtf.csrf import CSRFProtect
from flask_login import LoginManager, login_user, logout_user, login_required, current_user
from openpyxl.chart import DoughnutChart, Reference, BarChart, LineChart
from weasyprint import HTML
from io import BytesIO
import pandas as pd
import openpyxl
import matplotlib.pyplot as plt
import io
import base64

from config import config

# Models:
from models.ModelUser import ModelUser
from models.ModelGranja import ModelGranja
from models.ModelReporte import ModelReporte
from models.ModelVacuno import ModelVacuno
from models.ModelLote import ModelLote

# Entities:
from models.entities.User import User

app = Flask(__name__)

csrf = CSRFProtect(app)
db = MySQL(app)
login_manager_app = LoginManager(app)

def generate_doughnut(satisfactorio, no_satisfactorio):
    fig, ax = plt.subplots()
    sizes = [satisfactorio, no_satisfactorio]
    colors = ['#EAC655', '#D3D3D3']
    ax.pie(sizes, colors=colors, startangle=90, wedgeprops=dict(width=0.3))
    centre_circle = plt.Circle((0,0),0.70,fc='white')
    fig.gca().add_artist(centre_circle)
    buf = io.BytesIO()
    plt.savefig(buf, format='png', bbox_inches='tight')
    plt.close()
    return base64.b64encode(buf.getvalue()).decode('utf-8')


def generate_bar_chart(labels, values):
    plt.figure()
    bars = plt.bar(labels, values, color='#EAC655')
    plt.xticks(rotation=45)
    for bar in bars:
        plt.text(bar.get_x() + bar.get_width()/2, bar.get_height(),
                 int(bar.get_height()),
                 ha='center', va='bottom')
    buf = io.BytesIO()
    plt.savefig(buf, format='png', bbox_inches='tight')
    plt.close()
    return base64.b64encode(buf.getvalue()).decode('utf-8')

def generate_line_chart(labels, values):
    plt.figure()
    plt.plot(labels, values, marker='o', color='#C2A137')
    plt.xticks(rotation=45)
    plt.grid(True)
    buf = io.BytesIO()
    plt.savefig(buf, format='png', bbox_inches='tight')
    plt.close()
    return base64.b64encode(buf.getvalue()).decode('utf-8')

@app.template_filter('calcular_edad')
def calcular_edad(fecha_nacimiento):
    hoy = datetime.now().date()
    nacimiento = fecha_nacimiento
    edad = hoy.year - nacimiento.year
    #Revisa si el nacimiento ha ocurrido en este año
    if (hoy.month, hoy.day) < (nacimiento.month, nacimiento.day):
        edad -= 1
    return f"{edad} años"

@login_manager_app.user_loader
def load_user(id):
    return ModelUser.get_by_id(db, id)


@app.route('/')
def index():
    return redirect(url_for('login'))


@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        user = User(0, request.form['email'], request.form['password'])
        logged_user = ModelUser.login(db, user)
        if logged_user != None:
            if logged_user.password:
                login_user(logged_user)
                return redirect(url_for('home'))
            else:
                flash("Contraseña incorrecta...")
                return render_template('login.html')
        else:
            flash("El usuario no existe...")
            return render_template('login.html')
    else:
        return render_template('login.html')


@app.route('/logout')
def logout():
    logout_user()
    return redirect(url_for('login'))


@app.route('/home')
@login_required
def home():
    try:
        granjas = ModelGranja.get_by_owner(db, current_user.id)
        return render_template('InicioAdministrador.html', granjas=granjas)
    except Exception as ex:
        flash(f"Error loading farms: {str(ex)}")
        return redirect(url_for('login'))

@app.route('/recPassword')
def recPassword():
    return render_template('RecuperarContraseña.html')

@app.route('/estadisticas')
@login_required
def estadisticas():
    granja_id = request.args.get('granja_id')
    granja = ModelGranja.get_by_id(db, granja_id, current_user.id)
    if not granja:
        flash("Granja no encontrada")
        return redirect(url_for('home'))

    breed_stats = ModelGranja.get_breed_stats(db, granja_id)

    current_year = datetime.now().year
    monthly_stats = ModelReporte.get_monthly_cattle_stats(db, granja_id, current_year)

    monthly_counts = [0] * 12
    for month, total in monthly_stats:
        monthly_counts[month - 1] = total

    report_stats = ModelGranja.get_report_stats(db, granja_id)
    total = report_stats['with_reports'] + report_stats['without_reports']
    report_percent = round((report_stats['with_reports'] / total * 100)) if total > 0 else 0

    owner_name = current_user.nombre
    return render_template(
        'Estadisticas.html',
        granja=granja,
        owner_name=owner_name,
        monthly_counts=monthly_counts,
        razas=breed_stats,
        report_percent=report_percent,
        report_data=[
            report_stats['with_reports'],
            report_stats['without_reports']
        ]
    )

@app.route('/header')
def header():
    return render_template('header.html')

@app.route('/get_granjas')
@login_required
def get_granjas():
    try:
        granjas = ModelGranja.get_by_owner(db, current_user.id)
        return jsonify([{'id': g[0], 'direccion': g[1]} for g in granjas])
    except Exception as ex:
        return jsonify({'error': str(ex)}), 500

@app.route('/add_granja', methods=['POST'])
@login_required
def add_granja():
    try:
        data = request.json
        ModelGranja.create_granja(
            db=db,
            direccion=data['direccion'],
            cant_ganado=0,
            status=1,
            tatuaje=data['tatuaje'],
            id_dueño=current_user.id
        )
        return jsonify({'success': True})
    except ValueError as ve:
        return jsonify({'error': str(ve)}), 400
    except Exception as ex:
        app.logger.error(f"Error adding farm: {str(ex)}")
        return jsonify({'error': "Server error"}), 500

@app.route('/get_granja/<int:granja_id>')
@login_required
def get_granja(granja_id):
    try:
        granja = ModelGranja.get_by_id(db, granja_id, current_user.id)
        if not granja:
            return jsonify({'error': 'Granja no encontrada'}), 404

        return jsonify({
            'id': granja[0],
            'direccion': granja[1],
            'cant_ganado': granja[2],
            'status': granja[3],
            'tatuaje': granja[4]
        })
    except Exception as ex:
        return jsonify({'error': str(ex)}), 500

@app.route('/update_granja/<int:granja_id>', methods=['PUT'])
@login_required
def update_granja(granja_id):
    try:
        data = request.get_json()
        success = ModelGranja.update_granja(
            db=db,
            granja_id=granja_id,
            owner_id=current_user.id,
            direccion=data['direccion'],
            status=data['status'],
            tatuaje=data['tatuaje']
        )

        if not success:
            return jsonify({'error': 'Granja no encontrada o no autorizada'}), 404

        return jsonify({'success': True})
    except ValueError as ve:
        return jsonify({'error': str(ve)}), 400
    except Exception as ex:
        return jsonify({'error': str(ex)}), 500

@app.route('/delete_granja/<int:granja_id>', methods=['DELETE'])
@login_required
def delete_granja(granja_id):
    try:
        success = ModelGranja.delete_granja(db, granja_id, current_user.id)
        if not success:
            return jsonify({'error': 'Granja no encontrada o no autorizada'}), 404
        return jsonify({'success': True})
    except Exception as ex:
        return jsonify({'error': str(ex)}), 500


@app.route('/gestion_ganado')
@login_required
def gestion_ganado():
    granja_id = request.args.get('granja_id')
    granja = ModelGranja.get_by_id(db, granja_id, current_user.id)

    if not granja:
        flash("Granja no encontrada")
        return redirect(url_for('home'))

    vacunos = ModelVacuno.get_vacunos_by_granja(db, granja_id)
    lotes = ModelLote.get_active_lotes(db, granja_id)

    return render_template('GestionGanado.html',
                           vacunos=vacunos,
                           granja=granja,
                           lotes=lotes)


@app.route('/crear_vacuno', methods=['POST'])
@login_required
def crear_vacuno():
    try:
        data = request.get_json()

        # Validate lote belongs to user's farm
        lote = ModelLote.get_lote_by_id(db, data['lote'], current_user.id)
        if not lote:
            return jsonify({'success': False, 'error': 'Lote no válido'}), 400

        ModelVacuno.create_vacuno(db, {
            'lote': data['lote'],
            'etiqueta': data['etiqueta'],
            'fierro': data['fierro'],
            'nombre': data['nombre'],
            'sexo': data['sexo'],
            'fecha_nacimiento': data['fecha_nacimiento'],
            'raza': data['raza'],
            'proposito': data['proposito']
        })

        return jsonify({'success': True})

    except Exception as ex:
        app.logger.error(f"Error creating cattle: {str(ex)}")
        return jsonify({'success': False, 'error': str(ex)}), 500

@app.route('/actualizar_vacuno', methods=['POST'])
@login_required
def actualizar_vacuno():
    try:
        data = request.get_json()

        # Validate required fields
        required_fields = ['etiqueta', 'fierro', 'sexo', 'fecha_nacimiento', 'raza', 'proposito']
        if not all(field in data for field in required_fields):
            return jsonify({'success': False, 'error': 'Faltan campos requeridos'}), 400

        # Update in database
        updated = ModelVacuno.update_vacuno(
            db=db,
            etiqueta=data['etiqueta'],
            nuevos_datos={
                'fierro': data['fierro'],
                'sexo': data['sexo'],
                'fecha_nacimiento': data['fecha_nacimiento'],
                'raza': data['raza'],
                'proposito': data['proposito']
            }
        )

        if not updated:
            return jsonify({'success': False, 'error': 'Vacuno no encontrado'}), 404

        return jsonify({'success': True})

    except Exception as ex:
        app.logger.error(f"Error updating cattle: {str(ex)}")
        return jsonify({'success': False, 'error': str(ex)}), 500


@app.route('/eliminar_vacuno', methods=['POST'])
@login_required
def eliminar_vacuno():
    try:
        data = request.get_json()
        etiqueta = data.get('etiqueta')

        if not etiqueta:
            return jsonify({'success': False, 'error': 'Falta la etiqueta del vacuno'}), 400

        deleted = ModelVacuno.delete_vacuno(
            db=db,
            etiqueta=etiqueta,
            owner_id=current_user.id
        )

        if not deleted:
            return jsonify({'success': False, 'error': 'Vacuno no encontrado o no autorizado'}), 404

        return jsonify({'success': True})

    except Exception as ex:
        app.logger.error(f"Error deleting cattle: {str(ex)}")
        return jsonify({'success': False, 'error': str(ex)}), 500


@app.route('/exportar_pdf')
@login_required
def exportar_pdf():
    granja_id = request.args.get('granja_id')
    granja = ModelGranja.get_by_id(db, granja_id, current_user.id)

    if not granja:
        flash("Granja no encontrada")
        return redirect(url_for('home'))

    try:
        # Get all necessary data using existing models
        breed_stats = ModelGranja.get_breed_stats(db, granja_id)
        current_year = datetime.now().year
        monthly_stats = ModelReporte.get_monthly_cattle_stats(db, granja_id, current_year)
        report_stats = ModelGranja.get_report_stats(db, granja_id)

        # Process data for charts
        total_vacunos = report_stats['with_reports'] + report_stats['without_reports']
        report_percent = round((report_stats['with_reports'] / total_vacunos * 100)) if total_vacunos > 0 else 0

        # Generate chart images
        estado_chart = generate_doughnut(report_stats['with_reports'], report_stats['without_reports'])
        raza_chart = generate_bar_chart(list(breed_stats.keys()), list(breed_stats.values()))

        meses_labels = ['Ene', 'Feb', 'Mar', 'Abr', 'May', 'Jun',
                        'Jul', 'Ago', 'Sep', 'Oct', 'Nov', 'Dic']
        meses_values = [0] * 12
        for mes in monthly_stats:
            meses_values[mes[0] - 1] = mes[1]
        mes_chart = generate_line_chart(meses_labels, meses_values)

        html = render_template(
            'pdf_template.html',
            granja=granja,
            owner_name=current_user.nombre,
            estado_chart=estado_chart,
            raza_chart=raza_chart,
            mes_chart=mes_chart,
            report_percent=report_percent
        )

        pdf = HTML(string=html, base_url=request.host_url).write_pdf()
        response = make_response(pdf)
        response.headers['Content-Type'] = 'application/pdf'
        response.headers['Content-Disposition'] = f'attachment; filename=reporte_granja_{granja_id}.pdf'
        return response

    except Exception as ex:
        app.logger.error(f"Error generating PDF: {str(ex)}")
        flash("Error al generar el reporte PDF")
        return redirect(url_for('estadisticas', granja_id=granja_id))


@app.route('/exportar_excel')
@login_required
def exportar_excel():
    granja_id = request.args.get('granja_id')
    granja = ModelGranja.get_by_id(db, granja_id, current_user.id)

    if not granja:
        flash("Granja no encontrada")
        return redirect(url_for('home'))

    try:
        # Get all necessary data
        lotes = ModelReporte.get_lotes_for_export(db, granja_id)
        lote_ids = [lote[0] for lote in lotes]
        vacunos = ModelReporte.get_vacunos_for_export(db, lote_ids)
        reportes = ModelReporte.get_reportes_by_granja(db, granja_id)

        # Get chart data
        breed_stats = ModelGranja.get_breed_stats(db, granja_id)
        current_year = datetime.now().year
        monthly_stats = ModelReporte.get_monthly_cattle_stats(db, granja_id, current_year)
        report_stats = ModelGranja.get_report_stats(db, granja_id)

        # Create DataFrames
        df_granja = pd.DataFrame([{
            'ID': granja[0],
            'Dirección': granja[1],
            'Cantidad Ganado': granja[2],
            'Estado': 'Activo' if granja[3] else 'Inactivo',
            'Tatuaje': granja[4]
        }])

        df_lotes = pd.DataFrame([{
            'ID Lote': lote[0],
            'Fecha Registro': lote[1],
            'Nombre': lote[2],
            'Cantidad Vacunos': lote[3],
            'Estado': 'Activo' if lote[4] else 'Inactivo'
        } for lote in lotes])

        df_vacunos = pd.DataFrame([{
            'ID Vaca': vaca[0],
            'ID Lote': vaca[1],
            'ID Arete': vaca[2],
            'Género': vaca[3],
            'Fecha Nacimiento': vaca[4],
            'Raza': vaca[5],
            'Estado Salud': vaca[6],
            'Proposito': vaca[7]
        } for vaca in vacunos])

        df_reportes = pd.DataFrame([{
            'ID Reporte': reporte[0],
            'ID Arete': reporte[1],
            'Fecha Reporte': reporte[2],
            'Tipo Reporte': reporte[3],
            'Observaciones': reporte[4]
        } for reporte in reportes])

        # Create Excel file in memory
        output = BytesIO()
        with pd.ExcelWriter(output, engine='openpyxl') as writer:
            # Write data sheets
            df_granja.to_excel(writer, sheet_name='Granja', index=False)
            df_lotes.to_excel(writer, sheet_name='Lotes', index=False)
            df_vacunos.to_excel(writer, sheet_name='Vacunos', index=False)
            df_reportes.to_excel(writer, sheet_name='Reportes', index=False)

            # Get workbook and create charts sheet
            workbook = writer.book
            chartsheet = workbook.create_sheet('Gráficos')

            # Add chart data tables
            _add_chart_data(workbook, chartsheet, report_stats, breed_stats, monthly_stats)

            # Create charts
            _create_excel_charts(workbook, chartsheet, report_stats, breed_stats, monthly_stats)

        output.seek(0)

        return send_file(
            output,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            download_name=f"reporte_granja_{granja_id}.xlsx",
            as_attachment=True
        )

    except Exception as ex:
        app.logger.error(f"Error generating Excel: {str(ex)}")
        flash("Error al generar el reporte")
        return redirect(url_for('estadisticas', granja_id=granja_id))


def _add_chart_data(workbook, sheet, report_stats, breed_stats, monthly_stats):
    # Estado de la finca data
    sheet.append(['Estado', 'Cantidad'])
    sheet.append(['Con Reportes', report_stats['with_reports']])
    sheet.append(['Sin Reportes', report_stats['without_reports']])

    # Razas data
    sheet.append([])
    sheet.append([])
    sheet.append(['Raza', 'Cantidad'])
    for raza, cantidad in breed_stats.items():
        sheet.append([raza, cantidad])

    # Registros mensuales data
    sheet.append([])
    sheet.append([])
    sheet.append(['Mes', 'Cantidad'])
    meses = ['Ene', 'Feb', 'Mar', 'Abr', 'May', 'Jun',
             'Jul', 'Ago', 'Sep', 'Oct', 'Nov', 'Dic']

    # Process monthly stats into ordered list
    monthly_counts = [0] * 12
    for month, count in monthly_stats:
        if 1 <= month <= 12:
            monthly_counts[month - 1] = count

    for i, count in enumerate(monthly_counts):
        sheet.append([meses[i], count])


def _create_excel_charts(workbook, sheet, report_stats, breed_stats, monthly_stats):
    # Doughnut Chart - Estado de la Finca
    doughnut_chart = DoughnutChart()
    labels = Reference(sheet, min_col=1, min_row=2, max_row=3)
    data = Reference(sheet, min_col=2, min_row=1, max_row=3)
    doughnut_chart.add_data(data, titles_from_data=True)
    doughnut_chart.set_categories(labels)
    doughnut_chart.title = 'Estado de la Finca'
    doughnut_chart.style = 26
    sheet.add_chart(doughnut_chart, "E2")

    # Bar Chart - Distribución por Razas
    bar_chart = BarChart()
    bar_chart.title = "Distribución por Razas"
    bar_chart.style = 10
    bar_chart.y_axis.title = 'Cantidad'
    bar_chart.x_axis.title = 'Razas'

    data = Reference(sheet, min_col=2, min_row=6, max_row=6 + len(breed_stats))
    cats = Reference(sheet, min_col=1, min_row=6, max_row=6 + len(breed_stats))
    bar_chart.add_data(data)
    bar_chart.set_categories(cats)
    sheet.add_chart(bar_chart, "E20")

    # Line Chart - Registros Mensuales
    line_chart = LineChart()
    line_chart.title = "Registros Mensuales"
    line_chart.style = 12
    line_chart.y_axis.title = 'Cantidad'
    line_chart.x_axis.title = 'Meses'

    start_row = 10 + len(breed_stats)
    data = Reference(sheet, min_col=2, min_row=start_row, max_row=start_row + 11)
    cats = Reference(sheet, min_col=1, min_row=start_row, max_row=start_row + 11)
    line_chart.add_data(data)
    line_chart.set_categories(cats)
    sheet.add_chart(line_chart, "E40")

    # Apply color styling
    for chart in [doughnut_chart, bar_chart, line_chart]:
        for series in chart.series:
            series.graphicalProperties.line.solidFill = "C2A137"  # Border color
            series.graphicalProperties.solidFill = "EAC655"  # Fill color

@app.route('/importar_vacunos', methods=['POST'])
@login_required
def importar_vacunos():
    try:
        if 'file' not in request.files:
            return jsonify({'success': False, 'error': 'No file uploaded'}), 400

        file = request.files['file']
        if file.filename == '':
            return jsonify({'success': False, 'error': 'No selected file'}), 400

        if not file.filename.endswith('.xlsx'):
            return jsonify({'success': False, 'error': 'Solo se aceptan archivos XLSX'}), 400

        # Read Excel file
        df = pd.read_excel(file, dtype=str)
        required_columns = ['Etiqueta', 'Fierro', 'Sexo', 'Fecha_Nacimiento', 'Raza', 'Proposito', 'Lote_ID']

        # Validate columns
        if not all(col in df.columns for col in required_columns):
            return jsonify({
                'success': False,
                'error': f'El archivo debe contener estas columnas: {", ".join(required_columns)}'
            }), 400

        success_count = 0
        errors = []

        for index, row in df.iterrows():
            try:
                # Validate lote ownership
                lote = ModelLote.get_lote_by_id(db, row['Lote_ID'], current_user.id)
                if not lote:
                    errors.append(f"Fila {index + 2}: Lote ID {row['Lote_ID']} no válido")
                    continue

                # Prepare cattle data
                cattle_data = {
                    'lote': row['Lote_ID'],
                    'etiqueta': row['Etiqueta'],
                    'fierro': row['Fierro'],
                    'sexo': row['Sexo'],
                    'fecha_nacimiento': row['Fecha_Nacimiento'],
                    'raza': row['Raza'],
                    'proposito': row['Proposito']
                }

                # Insert into database
                ModelVacuno.create_vacuno(db, cattle_data)
                success_count += 1

            except Exception as e:
                errors.append(f"Fila {index + 2}: {str(e)}")
                continue

        return jsonify({
            'success': True,
            'imported': success_count,
            'errors': errors
        })

    except Exception as ex:
        return jsonify({
            'success': False,
            'error': f'Error procesando archivo: {str(ex)}'
        }), 500


@app.route('/gestion_reportes')
@login_required
def gestion_reportes():
    granja_id = request.args.get('granja_id')
    granja = ModelGranja.get_by_id(db, granja_id, current_user.id)

    if not granja:
        flash("Granja no encontrada")
        return redirect(url_for('home'))

    reportes = ModelReporte.get_reportes_by_granja(db, granja_id)
    return render_template('GestionReportes.html', granja=granja, reportes=reportes)


@app.route('/crear_reporte', methods=['POST'])
@login_required
def crear_reporte():
    try:
        data = request.get_json()
        granja_id = data.get('granja_id')

        # Validate farm ownership
        granja = ModelGranja.get_by_id(db, granja_id, current_user.id)
        if not granja:
            return jsonify({'success': False, 'error': 'Granja no válida'}), 400

        ModelReporte.create_reporte(db, {
            'etiqueta': data['etiqueta'],
            'fecha_reporte': data['fecha_reporte'],
            'tipo_reporte': data['tipo_reporte'],
            'observacion': data['observacion']
        }, granja_id)

        return jsonify({'success': True})

    except Exception as ex:
        return jsonify({'success': False, 'error': str(ex)}), 500


@app.route('/eliminar_reporte', methods=['POST'])
@login_required
def eliminar_reporte():
    try:
        data = request.get_json()
        reporte_id = data.get('reporte_id')

        if not reporte_id:
            return jsonify({'success': False, 'error': 'Falta el ID del reporte'}), 400

        deleted = ModelReporte.delete_reporte(
            db=db,
            reporte_id=reporte_id,
            owner_id=current_user.id
        )

        if not deleted:
            return jsonify({'success': False, 'error': 'Reporte no encontrado o no autorizado'}), 404

        return jsonify({'success': True})

    except Exception as ex:
        return jsonify({'success': False, 'error': str(ex)}), 500


@app.route('/informacion_vaca')
@login_required
def informacion_vaca():
    etiqueta = request.args.get('etiqueta')
    if not etiqueta:
        flash("Falta el parámetro de etiqueta")
        return redirect(url_for('home'))

    try:
        vacuno = ModelVacuno.get_vacuno_by_etiqueta(db, etiqueta, current_user.id)

        if not vacuno:
            flash("Vaca no encontrada o no tienes permisos")
            return redirect(url_for('home'))

        # Get reports
        reportes = ModelReporte.get_reportes_by_vaca(db, vacuno[0])  # [0] = ID_Vaca

        return render_template('InformacionVaca.html',
                               vacuno=vacuno,
                               reportes=reportes,
                               calcular_edad=calcular_edad)

    except Exception as ex:
        flash(f"Error al cargar datos: {str(ex)}")
        app.logger.error(f"Error en informacion_vaca: {str(ex)}")
        return redirect(url_for('gestion_ganado'))


@app.route('/gestion_lotes')
@login_required
def gestion_lotes():
    granja_id = request.args.get('granja_id')
    granja = ModelGranja.get_by_id(db, granja_id, current_user.id)

    if not granja:
        flash("Granja no encontrada")
        return redirect(url_for('home'))

    lotes = ModelLote.get_lotes_by_granja(db, granja_id)
    return render_template('GestionLotes.html',
                           granja=granja,
                           lotes=lotes)


@app.route('/crear_lote', methods=['POST'])
@login_required
def crear_lote():
    try:
        data = request.get_json()
        granja_id = data.get('granja_id')

        # Verify farm ownership
        granja = ModelGranja.get_by_id(db, granja_id, current_user.id)
        if not granja:
            return jsonify({'success': False, 'error': 'Granja no encontrada'}), 404

        # Create new lote
        lote_id = ModelLote.create_lote(
            db=db,
            nombre=data['nombre'],
            granja_id=granja_id
        )

        return jsonify({'success': True, 'lote_id': lote_id})

    except Exception as ex:
        return jsonify({'success': False, 'error': str(ex)}), 500


@app.route('/eliminar_lote', methods=['POST'])
@login_required
def eliminar_lote():
    try:
        data = request.get_json()
        lote_id = data.get('lote_id')

        deleted = ModelLote.delete_lote(
            db=db,
            lote_id=lote_id,
            owner_id=current_user.id
        )

        if not deleted:
            return jsonify({'success': False, 'error': 'Lote no encontrado'}), 404

        return jsonify({'success': True})

    except Exception as ex:
        return jsonify({'success': False, 'error': str(ex)}), 500


@app.route('/cambiar_estado_lote', methods=['POST'])
@login_required
def cambiar_estado_lote():
    try:
        data = request.get_json()
        lote_id = data.get('lote_id')

        success = ModelLote.toggle_status(
            db=db,
            lote_id=lote_id,
            owner_id=current_user.id
        )

        if not success:
            return jsonify({'success': False, 'error': 'Lote no encontrado'}), 404

        return jsonify({'success': True})

    except Exception as ex:
        return jsonify({'success': False, 'error': str(ex)}), 500

@app.route('/protected')
@login_required
def protected():
    return "<h1>Esta es una vista protegida, solo para usuarios autenticados.</h1>"


def status_401(error):
    return redirect(url_for('login'))


def status_404(error):
    return "<h1>Página no encontrada</h1>", 404


if __name__ == '__main__':
    app.config.from_object(config['development'])
    csrf.init_app(app)
    app.register_error_handler(401, status_401)
    app.register_error_handler(404, status_404)
    app.run()